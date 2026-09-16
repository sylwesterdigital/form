from __future__ import annotations

from pathlib import Path


def extract_cv_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
    elif suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        blocks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                blocks.append(" | ".join(cell.text for cell in row.cells))
        text = "\n".join(blocks)
    elif suffix in {".txt", ".md"}:
        text = path.read_text(encoding="utf-8", errors="replace")
    else:
        raise ValueError(f"Unsupported CV format: {suffix}")

    text = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if len(text) < 40:
        raise ValueError("The CV did not contain enough extractable text.")
    return text
